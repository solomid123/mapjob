# -*- coding: utf-8 -*-
"""
Read a CV, propose a profile.

The profile page asks for forty facts that are already written down, in order,
in a document the candidate has had for years. Typing them again is the reason
the page sits half empty, and a half-empty profile is what the tailoring step,
the form filler and the letter writer all draw their material from -- so the
cost of the typing is not boredom, it is a worse application.

So: upload the CV, and every field it states is read out of it. What comes back
is a *proposal*, not a save. Nothing here writes to the profile, because a
document reader that edits the record directly is one bad extraction away from
replacing a real employment history with a mangled one, and the page it feeds
is the only copy of anything typed by hand. The page fills its boxes with the
proposal and the person doing the applying decides what stays.

Three rules the mapping keeps:

  * Copy, never compose. Dates, job titles and achievements go across in the
    CV's own words and language. An extractor that improves a bullet has
    invented a claim the candidate will have to defend in an interview, and
    the tailoring step downstream is already built on the same rule.
  * Only the keys profile_store.EDITABLE allows, intersected with the fields a
    CV can actually state. The password keys sit in the same record; a model
    that echoed one back could never set it.
  * The file is read in memory and dropped. A CV is a home address, a phone
    number and an employment history in one place -- it is not written to disk
    here, and never to a public bucket.

With no model configured this still works, just narrowly: the contact details
are the part of a CV that is machine-readable without help.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Tuple

from services.automation.llm_client import FuelixClient
from services.automation import profile_store

# What a CV can honestly be said to state. The intersection with EDITABLE is
# taken here, so a key removed from the store's allow-list disappears from this
# module without anyone having to remember to.
FROM_A_CV: tuple = (
    "first_name", "last_name", "full_name", "email", "phone", "phone_formatted",
    "address", "postal_code", "city", "country", "full_address",
    "linkedin", "website", "github", "portfolio",
    "current_title", "headline", "summary", "years_of_experience",
    "experiences", "education", "technical_skills", "soft_skills", "languages",
    "certifications", "projects", "awards", "interests",
    "driving_licence", "nationality",
)

ALLOWED = tuple(k for k in FROM_A_CV
                if k in profile_store.EDITABLE and k not in profile_store.SECRET)

# Long enough for a four-page CV, short enough that a 200-page PDF renamed to
# look like one cannot turn into a very expensive prompt.
MAX_TEXT = 24000
MAX_SUMMARY = 2000
MAX_LINE = 400
MAX_ROLES = 20
MAX_SCHOOLS = 12
MAX_LIST = 80
MAX_CHIP = 120


# ---------------------------------------------------------------------------
# Getting the words out
# ---------------------------------------------------------------------------

def text_from_pdf(raw: bytes) -> Tuple[str, int]:
    """PyMuPDF where it is installed, pypdf otherwise. Both read a two-column
    CV badly enough that the structuring is left to the model rather than to a
    layout heuristic written here."""
    try:
        import fitz  # type: ignore

        with fitz.open(stream=raw, filetype="pdf") as doc:
            pages = [p.get_text("text") for p in doc]
        return "\n\n".join(pages), len(pages)
    except ImportError:
        pass
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [(p.extract_text() or "") for p in reader.pages]
    return "\n\n".join(pages), len(pages)


def text_from_docx(raw: bytes) -> str:
    import docx  # type: ignore

    document = docx.Document(io.BytesIO(raw))
    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(b for b in blocks if b.strip())


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace(" ", " ")
    # A designed CV puts a glyph in front of every bullet. It is noise in a
    # prompt and noise in an achievement chip.
    text = re.sub(r"[•●▪·]+[ \t]*", "- ", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# What can be read without a model
# ---------------------------------------------------------------------------

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]{2,}")
# Loose on purpose: +33 6 12 34 56 78, 06.12.34.56.78, (0) 171 234 5678.
PHONE_RE = re.compile(r"(?:\+\d{1,3}[\s.\-/]?)?(?:\(0\)[\s.\-]?)?\d(?:[\s.\-/]?\d){7,13}")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:[\w-]+\.)?linkedin\.com/[\w%/\-.]+", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w\-.]+", re.I)
SITE_RE = re.compile(r"(?:https?://)?(?:www\.)?[\w-]+\.[a-z]{2,10}(?:/[\w%\-./]*)?", re.I)
NOT_A_SITE = ("linkedin.com", "github.com", "gmail.com", "outlook.com",
              "hotmail.com", "yahoo.", "orange.fr", "free.fr", "icloud.com",
              "wa.me", "t.me")


def _first(pattern, text: str) -> str:
    found = pattern.search(text)
    return found.group(0).strip() if found else ""


def patterns(text: str) -> Dict[str, Any]:
    """The contact block, which is the same shape on every CV ever printed."""
    head = text[:1800]
    out: Dict[str, Any] = {}

    email = _first(EMAIL_RE, text)
    if email:
        out["email"] = email.rstrip(".,;")

    # Phone numbers are looked for in the header only: further down the page a
    # date range, a postcode and a reference number all match a loose pattern.
    for candidate in PHONE_RE.findall(head):
        digits = re.sub(r"\D", "", candidate)
        if 9 <= len(digits) <= 15:
            out["phone_formatted"] = candidate.strip()
            break

    link = _first(LINKEDIN_RE, text)
    if link:
        out["linkedin"] = link if link.startswith("http") else "https://" + link
    hub = _first(GITHUB_RE, text)
    if hub:
        out["github"] = hub if hub.startswith("http") else "https://" + hub

    # The addresses come out first: the tail of an email is a perfectly good
    # domain, and "example.com" out of someone@example.com is not a website.
    for candidate in SITE_RE.findall(EMAIL_RE.sub(" ", head)):
        low = candidate.lower()
        if any(bad in low for bad in NOT_A_SITE) or "@" in candidate:
            continue
        if low.endswith((".pdf", ".png", ".jpg")):
            continue
        out["website"] = candidate if candidate.startswith("http") else "https://" + candidate
        break

    return out


# ---------------------------------------------------------------------------
# The mapping
# ---------------------------------------------------------------------------

SYSTEM = ("You transcribe a CV into structured fields. You copy what the "
          "document says and you never add anything to it.")

SCHEMA = """{
  "first_name": "", "last_name": "", "full_name": "",
  "email": "", "phone_formatted": "",
  "address": "street and number", "postal_code": "", "city": "", "country": "",
  "linkedin": "", "website": "", "github": "",
  "current_title": "the job title of the most recent role",
  "headline": "the line printed under the name, if there is one",
  "summary": "the profile or about paragraph, word for word, if there is one",
  "years_of_experience": 0,
  "experiences": [
    {"role": "", "company": "", "location": "",
     "period": "as printed, e.g. Aug 2025 - Present",
     "highlights": ["one bullet, as written"]}
  ],
  "education": [
    {"degree": "", "institution": "", "location": "", "period": "", "note": ""}
  ],
  "technical_skills": [], "soft_skills": [], "certifications": [],
  "languages": {"French": "native", "English": "C1"},
  "projects": [], "awards": [], "interests": [],
  "driving_licence": "", "nationality": ""
}"""

RULES = (
    "Fill the JSON below from the CV text that follows it.\n"
    "\n"
    "- Copy. Do not rewrite, translate, summarise, expand or improve anything. "
    "Bullets keep their own wording and their own language. Dates keep the form "
    "they are printed in.\n"
    "- A field the CV does not state is left out of your answer entirely. Do not "
    "guess a country from a city, a seniority from a date, or a skill from a job "
    "title. A missing field is correct; an invented one is a claim the candidate "
    "has to defend in an interview.\n"
    "- years_of_experience: only if the CV states a number of years. Never "
    "compute it from dates.\n"
    "- experiences: most recent first, every role the CV lists.\n"
    "- languages: an object of language to level, in the CV's own words.\n"
    "- Return the JSON object and nothing else. No commentary, no code fences, "
    "no placeholders in square brackets.\n"
)


def _text(value: Any, cap: int = MAX_LINE) -> str:
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    if not isinstance(value, str):
        return ""
    out = re.sub(r"\s+", " ", value).strip()
    # A model that could not find a field sometimes writes down the fact that
    # it could not find it. That is not a value.
    if out.lower() in ("n/a", "na", "none", "null", "unknown", "not stated",
                       "not specified", "-", "--"):
        return ""
    if out.startswith("[") and out.endswith("]"):
        return ""
    return out[:cap]


def _chips(value: Any) -> List[str]:
    if isinstance(value, str):
        value = re.split(r"[,;|]", value)
    if not isinstance(value, (list, tuple)):
        return []
    out: List[str] = []
    for item in value:
        if isinstance(item, dict):
            item = item.get("name") or item.get("title") or ""
        chip = _text(item, MAX_CHIP)
        if chip and chip not in out:
            out.append(chip)
    return out[:MAX_LIST]


def _bullets(item: Dict[str, Any]) -> List[str]:
    given = item.get("highlights") or item.get("bullets") or item.get("achievements")
    if isinstance(given, str):
        given = [line for line in given.split("\n")]
    if not isinstance(given, (list, tuple)):
        return []
    out = [_text(x, MAX_LINE) for x in given]
    return [x for x in out if x][:12]


def _roles(value: Any) -> List[Dict[str, Any]]:
    if not isinstance(value, (list, tuple)):
        return []
    out: List[Dict[str, Any]] = []
    for item in value[:MAX_ROLES]:
        if not isinstance(item, dict):
            continue
        role = {
            "role": _text(item.get("role") or item.get("title")),
            "company": _text(item.get("company") or item.get("employer")),
            "location": _text(item.get("location")),
            "period": _text(item.get("period") or item.get("dates")),
            "highlights": _bullets(item),
        }
        if role["role"] or role["company"]:
            out.append(role)
    return out


def _schools(value: Any) -> List[Dict[str, Any]]:
    if not isinstance(value, (list, tuple)):
        return []
    out: List[Dict[str, Any]] = []
    for item in value[:MAX_SCHOOLS]:
        if not isinstance(item, dict):
            continue
        school = {
            "degree": _text(item.get("degree") or item.get("title")),
            "institution": _text(item.get("institution") or item.get("school")),
            "location": _text(item.get("location")),
            "period": _text(item.get("period") or item.get("dates")),
            "note": _text(item.get("note") or item.get("details")),
        }
        if school["degree"] or school["institution"]:
            out.append(school)
    return out


def _levels(value: Any) -> Dict[str, str]:
    out: Dict[str, str] = {}
    pairs: List[Tuple[Any, Any]] = []
    if isinstance(value, dict):
        pairs = list(value.items())
    elif isinstance(value, (list, tuple)):
        for item in value:
            if isinstance(item, dict):
                pairs.append((item.get("language") or item.get("name") or "",
                              item.get("level") or item.get("proficiency") or ""))
            elif isinstance(item, str) and ":" in item:
                name, _, level = item.partition(":")
                pairs.append((name, level))
            elif isinstance(item, str):
                pairs.append((item, ""))
    for name, level in pairs:
        key = _text(name, 60)
        if key:
            out[key] = _text(level, 60)
    return dict(list(out.items())[:15])


def clean(raw: Any) -> Tuple[Dict[str, Any], List[str]]:
    """
    A model's answer, reduced to fields this app is willing to hold.

    Returns the fields and the names of everything refused, because a key that
    was quietly dropped is a key somebody spends an afternoon looking for.
    """
    if not isinstance(raw, dict):
        return {}, []
    fields: Dict[str, Any] = {}
    refused: List[str] = []

    for key, value in raw.items():
        if key not in ALLOWED:
            refused.append(str(key)[:40])
            continue
        if key == "experiences":
            got: Any = _roles(value)
        elif key == "education":
            got = _schools(value)
        elif key == "languages":
            got = _levels(value)
        elif key in ("technical_skills", "soft_skills", "certifications",
                     "projects", "awards", "interests"):
            got = _chips(value)
        elif key == "years_of_experience":
            try:
                number = float(str(value).replace(",", "."))
            except (TypeError, ValueError):
                continue
            if not 0 < number <= 60:
                continue
            got = int(number) if number == int(number) else round(number, 1)
        elif key == "summary":
            got = _text(value, MAX_SUMMARY)
        else:
            got = _text(value)
        if got not in ("", [], {}, None):
            fields[key] = got

    # A CV prints "linkedin.com/in/someone" because a printed page has nothing
    # to click. The forms this fills want a link.
    for key in ("linkedin", "website", "github", "portfolio"):
        value = fields.get(key)
        if isinstance(value, str) and value and not value.startswith(("http://", "https://")):
            fields[key] = "https://" + value.lstrip("/")

    # A CV prints the address over three lines; forms with a single address box
    # want it over one. Assembled from the parts just read, so it cannot say
    # something the other four fields do not -- and never derived when the
    # street or the city is missing, because half an address on a form is worse
    # than none.
    if not fields.get("full_address") and fields.get("address") and fields.get("city"):
        line = fields["address"] + ", "
        if fields.get("postal_code"):
            line += fields["postal_code"] + " "
        line += fields["city"]
        if fields.get("country"):
            line += ", " + fields["country"]
        fields["full_address"] = line

    # Two names and no full name, or the reverse. This is arithmetic on what
    # the CV already said, not a fact from nowhere.
    if fields.get("first_name") and fields.get("last_name") and not fields.get("full_name"):
        fields["full_name"] = fields["first_name"] + " " + fields["last_name"]
    if fields.get("full_name") and not (fields.get("first_name") or fields.get("last_name")):
        parts = fields["full_name"].split()
        if len(parts) >= 2:
            fields["first_name"] = parts[0]
            fields["last_name"] = " ".join(parts[1:])

    return fields, sorted(set(refused))


def read(raw: bytes, filename: str = "") -> Dict[str, Any]:
    """
    The proposal: {fields, found, refused, source, pages, chars, notes}.

    It does not raise on a file that simply gave up little. An empty `fields`
    with a note saying why is a better answer than an error, because the page
    behind it still works by hand.
    """
    name = (filename or "CV").strip()
    notes: List[str] = []
    refused: List[str] = []

    if name.lower().endswith(".docx"):
        text, pages = text_from_docx(raw), 0
    else:
        text, pages = text_from_pdf(raw)
    text = tidy(text)

    if not text:
        return {"filename": name, "pages": pages, "chars": 0, "fields": {},
                "found": [], "refused": [], "source": "none",
                "notes": ["There is no text in that file. A scanned CV is a "
                          "picture of words, and reading pictures is not "
                          "something this does."]}

    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT]
        notes.append("Only the first " + str(MAX_TEXT) + " characters were read.")

    fields = patterns(text)
    source = "patterns"

    try:
        client = FuelixClient()
        if not client.has_credentials():
            notes.append("No language model is configured, so only the contact "
                         "details could be read.")
        else:
            answer = client.chat_json(
                SYSTEM, RULES + "\nJSON SHAPE:\n" + SCHEMA + "\n\nCV TEXT:\n" + text,
                task="planner", temperature=0.0, max_tokens=4000, timeout=90)
            mapped, refused = clean(answer)
            if mapped:
                # The patterns keep whatever the model missed and lose to it on
                # anything it found: it saw the page, they saw a string.
                fields = {**fields, **mapped}
                source = "model"
            else:
                notes.append("The reader could not make sense of the layout, so "
                             "only the contact details came across.")
    except Exception as error:  # noqa: BLE001 - a partial fill beats a failure
        notes.append("The reader failed (" + str(error)[:120] + "). The contact "
                     "details were still read.")

    fields = {k: v for k, v in fields.items() if k in ALLOWED}
    return {
        "filename": name,
        "pages": pages,
        "chars": len(text),
        "source": source,
        "fields": fields,
        "found": sorted(fields.keys()),
        "refused": refused,
        "notes": notes,
    }
